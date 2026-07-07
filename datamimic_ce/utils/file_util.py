# DATAMIMIC
# Copyright (c) 2023-2025 Rapiddweller Asia Co., Ltd.
# This software is licensed under the MIT License.
# See LICENSE file for the full text of the license.
# For questions and support, contact: info@rapiddweller.com

import csv
import json
import shutil
import zipfile
from pathlib import Path

import numpy as np
import pandas as pd
from pandas import DataFrame

from datamimic_ce.utils.file_content_storage import FileContentStorage


class FileUtil:
    @staticmethod
    def read_dbunit_to_dict_list(path: Path, table: str) -> list[dict[str, str | None]]:
        """Read one table from a dbunit flat-XML dataset.

        Flat XML: every child of <dataset> is a row, the element name is the table, its attributes are
        the columns. A dataset holds many tables, so `table` selects one. Rows of one table may carry
        different columns ("ragged") - an absent attribute is a NULL, a present empty string is "". The
        result unifies the columns across the table (column sensing, like DbUnit >= 2.3): every row
        carries every column, an absent one filled with None. This gives faithful table semantics (a
        NULL cell, not a missing key) and lets a batch RDBMS insert see a uniform column set.

        Security: a value with XML character entities is decoded. ElementTree does NOT fetch external
        DTDs, but it DOES expand internal entities - dbunit files are trusted local fixtures, not
        untrusted input; do not point this at attacker-controlled XML.
        """
        import xml.etree.ElementTree as ET  # noqa: N817

        try:
            root = ET.parse(str(path)).getroot()
        except ET.ParseError as e:
            raise ValueError(f"dbunit dataset '{path}' is not well-formed XML: {e}") from e
        # dbunit's root is <dataset>; anything else is not a dataset
        if root.tag != "dataset":
            raise ValueError(f"dbunit dataset '{path}' must have a <dataset> root, got <{root.tag}>")
        raw = [child.attrib for child in root if child.tag == table]
        if not raw:
            available = sorted({child.tag for child in root})
            raise ValueError(f"dbunit dataset '{path}' has no rows for table '{table}'; available: {available}")
        # column sensing: union of columns in first-appearance order, absent -> None
        columns: dict[str, None] = {}
        for attrib in raw:
            columns.update(dict.fromkeys(attrib))
        return [{col: attrib.get(col) for col in columns} for attrib in raw]

    @staticmethod
    def parse_properties(path: Path, encoding="utf-8") -> dict[str, str]:
        """
        Parse properties from file then save into a dict
        :param path:
        :param encoding:
        :return:
        """

        try:
            properties_dict = {}
            # Load file content from cache or file
            data = FileContentStorage.load_file_with_custom_func(
                str(path), lambda: [line.strip() for line in path.open("r", encoding=encoding)]
            )
            # Parse properties from file content
            for line in data:
                # Skip comments and empty lines
                if line.strip() and not line.startswith("#"):
                    key, value = line.strip().split("=", 1)
                    properties_dict[key.strip()] = value.strip()
        except FileNotFoundError as e:
            raise FileNotFoundError(
                f"Property file not found {str(path)}, please check the file path again. Error message: {e}"
            ) from e

        return properties_dict

    @staticmethod
    def _read_raw_csv(file_path: Path, separator: str, encoding="utf-8") -> list[tuple]:
        """
        Read raw csv data
        """
        try:
            return FileContentStorage.load_file_with_custom_func(
                str(file_path),
                lambda: [
                    tuple(row)
                    for row in csv.reader(file_path.open("r", newline="", encoding=encoding), delimiter=separator)
                ],
            )
        except FileNotFoundError as e:
            raise FileNotFoundError(f"CSV file not found '{file_path}', error: {e}") from e

    @staticmethod
    def read_csv_to_dict_list(file_path: Path, separator: str, encoding="utf-8") -> list[dict]:
        """
        Read data from csv and parse into list of dict
        """
        raw_data = FileUtil._read_raw_csv(file_path, separator, encoding)
        if not raw_data:
            return []  # an empty CSV is an empty source, not a crash
        # Column names never carry meaningful surrounding whitespace; a padded/aligned CSV
        # (e.g. Benerator entity CSVs: "ean_code     ,name    ,...") would otherwise produce
        # keys like "name    " that a script's field access ("this.name") cannot resolve.
        header = [col.strip() if isinstance(col, str) else col for col in raw_data[0]]
        processed_data = [dict(zip(header, row, strict=False)) for row in raw_data[1:]]
        return processed_data

    @staticmethod
    def read_xlsx_to_dict_list(file_path: Path, sheet_name: str | None = None) -> list[dict]:
        """Read the first row of an .xlsx sheet as the header and each following row as a dict.

        Robust to real-world sheets: an empty sheet/file yields []; blank header cells are not turned
        into ``None``-keyed columns; a row shorter than the header pads missing cells with ``None`` and
        cells past the last header column are ignored. A file that is not a valid .xlsx raises a clear
        ValueError rather than a bare BadZipFile.
        """
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException

        try:
            workbook = load_workbook(file_path, read_only=True, data_only=True)
        except (InvalidFileException, zipfile.BadZipFile, KeyError) as e:
            raise ValueError(f"Invalid XLSX file '{file_path}': {e}") from e

        try:
            sheet = workbook[sheet_name] if sheet_name else workbook.active
        except KeyError as e:
            raise ValueError(f"XLSX file '{file_path}' has no sheet named '{sheet_name}'") from e
        if sheet is None:
            return []  # no sheet -> empty source

        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            return []  # empty sheet is an empty source, not a crash
        # Real columns = non-blank header cells, keyed by their column position (skip blank headers so
        # openpyxl's rectangular row padding never produces a None-keyed column).
        columns = [(idx, str(name)) for idx, name in enumerate(header) if name is not None]
        return [{name: (row[idx] if idx < len(row) else None) for idx, name in columns} for row in rows]

    @staticmethod
    def _parses_as_float(value: str) -> bool:
        try:
            float(value)
            return True
        except (TypeError, ValueError):
            return False

    @staticmethod
    def parse_fixed_width_spec(spec: str) -> list[tuple[str, int, bool, str]]:
        """Parse a Benerator-style fixed-width column spec: ``name[width]`` (left-aligned,
        space-padded) or ``name[width r pad]`` (right-aligned, e.g. ``price[8r0]`` = width 8,
        zero-padded). Returns ``(name, width, right_aligned, pad_char)`` per column, in order.
        """
        import re

        fields = []
        for token in spec.split(","):
            token = token.strip()
            match = re.fullmatch(r"(\w+)\[(\d+)(r)?(.)?\]", token)
            if not match:
                raise ValueError(f"Invalid fixed-width column spec token: '{token}' in '{spec}'")
            name, width, right_flag, pad = match.groups()
            right_aligned = right_flag is not None
            pad_char = pad if pad is not None else ("0" if right_aligned else " ")
            fields.append((name, int(width), right_aligned, pad_char))
        return fields

    @staticmethod
    def read_fixed_width_to_dict_list(file_path: Path, spec: str | None = None) -> list[dict]:
        """Read a fixed-width-column file into a list of dicts.

        Self-describing by default (``spec=None``): the file's first line must be
        ``# name[13],name2[30],...`` (the column spec as a comment) so the reader needs only the
        path - the same convention ``FixedWidthExporter`` writes. Pass ``spec`` explicitly to read
        a file that doesn't carry that header line.
        """
        lines = FileContentStorage.load_file_with_custom_func(
            str(file_path), lambda: file_path.read_text(encoding="utf-8").splitlines()
        )
        if not lines:
            return []  # an empty file is an empty source, not a crash

        if spec is not None:
            fields = FileUtil.parse_fixed_width_spec(spec)
            data_lines = lines
        else:
            header = lines[0]
            if not header.startswith("#"):
                raise ValueError(
                    f"Fixed-width file '{file_path}' has no '# name[width],...' spec header on its "
                    f"first line - pass spec= explicitly to read a file without one"
                )
            fields = FileUtil.parse_fixed_width_spec(header[1:])
            data_lines = lines[1:]

        result = []
        for line in data_lines:
            row = {}
            offset = 0
            for name, width, right_aligned, pad_char in fields:
                raw = line[offset : offset + width]
                row[name] = raw.lstrip(pad_char) if right_aligned else raw.strip()
                offset += width
            result.append(row)
        return result

    @staticmethod
    def read_weight_csv(file_path: Path, separator: str = ",", encoding="utf-8") -> DataFrame:
        """
        Read a 2-column value|weight csv, header optional. Auto-detected: if the first row's
        weight column doesn't parse as a number, it's a header row and gets skipped.
        """
        # Load file content from cache or file
        raw_data = FileUtil._read_raw_csv(file_path, separator, encoding)

        if raw_data and len(raw_data[0]) > 1 and not FileUtil._parses_as_float(raw_data[0][1]):
            raw_data = raw_data[1:]

        # Convert data to DataFrame, select only 2 columns (data and weight)
        df = pd.DataFrame(raw_data, columns=[0, 1])
        # Convert column 1 to float and replace NaN with 1
        df[1] = df[1].astype(float).fillna(1)
        # Replace NaN df in the first column with None to avoid nan
        df[0] = df[0].replace(to_replace=np.nan, value=None)
        # Calculate probability using count stat
        df[1] = df[1] / df[1].sum()

        return df

    @staticmethod
    def read_json(file_path: Path, encoding="utf-8") -> list[dict] | dict:
        """
        Read data from JSON
        """
        try:
            return FileContentStorage.load_file_with_custom_func(
                str(file_path), lambda: json.load(file_path.open(mode="r", encoding=encoding))
            )
        except FileNotFoundError as e:
            raise FileNotFoundError(f"JSON file not found '{file_path}', error: {e}") from e

    @staticmethod
    def read_json_to_list(file_path: Path, encoding="utf-8") -> list[dict]:
        """
        Read data from JSON and parse into list of dict
        """
        json_data = FileUtil.read_json(file_path, encoding)
        if isinstance(json_data, list):
            return json_data
        else:
            raise ValueError(f"JSON file '{file_path}' must contain a list of objects")

    @staticmethod
    def read_json_to_dict(file_path: Path, encoding="utf-8") -> dict:
        """
        Read data from JSON and parse into dict
        """
        json_data = FileUtil.read_json(file_path, encoding)
        if isinstance(json_data, dict):
            return json_data
        else:
            raise ValueError(f"JSON file '{file_path}' must contain a dictionary")

    @staticmethod
    def read_csv_to_dict_of_tuples_with_header(
        file_path: Path, delimiter: str = ",", encoding="utf-8"
    ) -> tuple[dict, list[tuple]]:
        """
        Read CSV to header dict and data list
        :param delimiter: delimiter used in the CSV file
        :param file_path: path to the CSV file
        :param encoding: encoding of the CSV file
        :return: a tuple containing a dictionary of headers and a list of tuples with string datas
        """
        # Load raw data
        raw_data = FileUtil._read_raw_csv(file_path, delimiter, encoding)
        header = raw_data[0]
        data_rows = raw_data[1:]

        # Create header dict
        header_dict = {}
        for idx, column in enumerate(header):
            modified_column = column.replace("\ufeff", "")
            header_dict[modified_column] = idx

        # Return header dict and data rows
        return header_dict, data_rows

    @staticmethod
    def read_csv_to_list_of_tuples_without_header(
        file_path: Path, delimiter: str = ",", encoding="utf-8"
    ) -> list[tuple]:
        """
        Read CSV without header to data list
        :param file_path: path to the CSV file
        :param delimiter: delimiter used in the CSV file
        :param encoding: encoding of the CSV file
        :return: a list of tuples containing the string data
        """
        return FileUtil._read_raw_csv(file_path, delimiter, encoding)

    @staticmethod
    def read_wgt_file(file_path: Path, delimiter: str = ",", encoding="utf-8") -> tuple[list, list]:
        """
        Read wgt file having no header and 2 columns (wgt is 2nd column)
        :param file_path:
        :param delimiter:
        :param encoding: encoding of the CSV file
        :return: Tuple contain list of values and list of weights
        """
        values = []
        weights = []

        # Load raw data
        raw_data = FileUtil._read_raw_csv(file_path, delimiter, encoding)

        # Process data
        for row in raw_data:
            #  tolerate text values containing delimiters; use last column as weight
            if len(row) >= 2:
                try:
                    weights.append(float(row[-1]))
                except (TypeError, ValueError):
                    # Treat non-numeric weight as 1.0
                    weights.append(1.0)
                # Reconstruct value by joining all but the last column
                values.append((",".join(row[:-1])).strip())
            elif len(row) == 1:
                # Assume weight as 1 if missing
                weights.append(1.0)
                values.append(row[0])
            else:
                # Skip empty rows
                continue

        # Normalize weights
        weights_sum = sum(weights)
        weights = [weight / weights_sum for weight in weights]

        return values, weights

    @staticmethod
    def read_csv_having_weight_column(filepath: Path, weight_column_name: str, delimiter: str = ",", encoding="utf-8"):
        """
        Read CSV file having one weight column
        :param filepath:
        :param weight_column_name:
        :param delimiter:
        :param encoding:
        :return:
        """
        weights = []  # List to store weights
        data_without_weights = []  # List to store dictionaries of data without the specified weight column

        # Load raw data
        list_of_dict_data = FileUtil.read_csv_to_dict_list(filepath, delimiter, encoding)

        # Process data
        for row in list_of_dict_data:
            # Extract and remove the specified weight column from the row
            weight = row.pop(weight_column_name, None)
            if weight is not None:
                # Convert weight to the appropriate type (float, int) if necessary
                weights.append(float(weight))
                # Add the modified row (now without the weight) to the data_without_weights list
                data_without_weights.append(row)

        # Return the tuple of weights list and data_without_weights list
        return (weights, data_without_weights)

    @staticmethod
    def read_mutil_column_wgt_file(
        file_path: Path,
        weight_col_index: int = 1,
        delimiter: str = ",",
        encoding="utf-8",
    ) -> tuple[list, list]:
        """
        Read wgt file having no header and mutil columns,
        if weight column missing or wrong index then weight value will be 1.0
        :param file_path:
        :param weight_col_index: index of weight column (default = 1)
        :param delimiter:
        :param encoding:
        :return: Tuple contain list of values and list of weights
        """
        weights = []  # List to store weights
        values = []  # List to store data

        # Load raw data
        raw_data = FileUtil._read_raw_csv(file_path, delimiter, encoding)

        for row in raw_data:
            # Skip the empty row
            if not row:
                continue
            # default weight 1.0 when weight column is missing or wrong index
            if weight_col_index < 0 or weight_col_index >= len(row):
                weights.append(1.0)
            else:
                weight = row[weight_col_index]
                weights.append(float(weight) if weight else 1.0)

            values.append(row)
        # Return the tuple of values list and weights list
        return values, weights

    @staticmethod
    def copy_file(source: Path, destination: Path) -> None:
        """
        Copy a file from source to destination.

        Args:
            source (Path): Source file path
            destination (Path): Destination file path
        """
        shutil.copy2(source, destination)

    @staticmethod
    def create_project_structure(project_dir: Path) -> None:
        """
        Create the initial project structure with necessary files and directories.

        Args:
            project_dir (Path): Target project directory
        """

        initial_descriptor_content = """
<setup>
    <generate name="datamimic_user_list" count="1000" target="CSV,JSON">
        <variable name="person" entity="Person(min_age=18, max_age=90, female_quota=0.5)"/>
        <key name="id" generator="IncrementGenerator"/>
        <key name="given_name" script="person.given_name"/>
        <key name="family_name" script="person.family_name"/>
        <key name="gender" script="person.gender"/>
        <key name="birthDate" script="person.birthdate" converter="DateFormat('%d.%m.%Y')"/>
        <key name="email" script="person.family_name + '@' + person.given_name + '.de'"/>
        <key name="ce_user" values="True, False"/>
        <key name="ee_user" values="True, False"/>
        <key name="datamimic_lover" constant="DEFINITELY"/>
    </generate>
</setup>
        """

        # Create basic directory structure
        (project_dir / "data").mkdir(exist_ok=True)
        (project_dir / "output").mkdir(exist_ok=True)
        (project_dir / "script").mkdir(exist_ok=True)
        (project_dir / "config").mkdir(exist_ok=True)

        # Create the descriptor file with the specified content
        descriptor_path = project_dir / "datamimic.xml"
        descriptor_path.write_text(initial_descriptor_content)

        # Create a default README.md
        readme_content = f"""
# DATAMIMIC Project: {project_dir.name}
This project was created using DATAMIMIC.

## Project Structure
- `data/`: Directory for input data files, like .ent.csv or .wgt.csv
- `script/`: Directory for input scripts or custom functions
- `output/`: Directory for generated output
- `config/`: Configuration files
- `datamimic.xml`: Main project descriptor file

## Initial Setup
The project is initialized with a sample descriptor that generates user data with the following fields:
- User ID (auto-incrementing)
- First Name
- Last Name
- Gender
- Birth Date
- Email
- CE User status
- EE User status
- DATAMIMIC Lover status

## Getting Started
1. Review and modify the `datamimic.xml` file to customize your data generation
2. Place any required input files in the `data/` directory
3. Run the project using: `datamimic run datamimic.xml`
        """
        (project_dir / "README.md").write_text(readme_content)
