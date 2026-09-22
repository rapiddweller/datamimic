# DATAMIMIC
# Copyright (c) 2023-2025 Rapiddweller Asia Co., Ltd.
# This software is licensed under the MIT License.
# See LICENSE file for the full text of the license.
# For questions and support, contact: info@rapiddweller.com

"""Exporter/source matrix. Self-extending gates:

* target x encoding: ONE generate writes non-ASCII rows to EVERY registered buffered exporter,
  once with the default (UTF-8) and once with cp1252; each file is decoded with the encoding it
  was written in and read back. A new exporter without a reader fails the coverage assert.
* source: every exportable format is read back through ``source=`` (plus committed weighted-CSV
  fixtures). A new ``SourceFileFormat`` without a case fails the coverage assert.
"""

import csv
import io
import json
import shutil
import xml.etree.ElementTree as ET
from pathlib import Path

import pytest

from datamimic_ce.data_mimic_test import DataMimicTest
from datamimic_ce.engine.dsl.constants.exporter_constants import (
    EXPORTER_CSV,
    EXPORTER_DBUNIT,
    EXPORTER_FIXED_WIDTH,
    EXPORTER_JSON,
    EXPORTER_TXT,
    EXPORTER_XLSX,
    EXPORTER_XML,
)
from datamimic_ce.engine.dsl.model.constraints import SourceFileFormat
from datamimic_ce.engine.io.exporters.exporter_util import _BUFFERED_EXPORTERS
from datamimic_ce.utils.file_util import FileUtil

_DIR = Path(__file__).resolve().parent
_OUT = _DIR / "output" / "matrix_out"
_COUNT = 10
_EXPECTED = [(i, f"Müller-€-{i}") for i in range(1, _COUNT + 1)]

Row = tuple[int, str]


def _file(ext: str) -> Path:
    # exact name: "*.xml" would also match the ".dbunit.xml" dataset
    hits = list(_OUT.rglob(f"matrix_rows.{ext}"))
    assert len(hits) == 1, f"expected exactly one matrix_rows.{ext} in {_OUT}, found {[p.name for p in hits]}"
    return hits[0]


def _read_csv(encoding: str) -> list[Row]:
    rows = csv.DictReader(io.StringIO(_file("csv").read_text(encoding=encoding), newline=""))
    return [(int(row["id"]), row["name"]) for row in rows]


def _read_json(encoding: str) -> list[Row]:
    return [(int(row["id"]), row["name"]) for row in json.loads(_file("json").read_text(encoding=encoding))]


def _read_xml(encoding: str) -> list[Row]:
    # Parsed from bytes on purpose: the XML declaration, not the test, must supply the encoding.
    root = ET.fromstring(_file("xml").read_bytes())
    return [(int(item.findtext("id")), item.findtext("name")) for item in root.iter("item")]


def _read_xlsx(encoding: str) -> list[Row]:
    from openpyxl import load_workbook

    sheet = load_workbook(_file("xlsx")).active
    rows = list(sheet.iter_rows(values_only=True))
    header, data = rows[0], rows[1:]
    id_col, name_col = header.index("id"), header.index("name")
    return [(int(row[id_col]), row[name_col]) for row in data]


def _read_txt(encoding: str) -> list[Row]:
    lines = [ln for ln in _file("txt").read_text(encoding=encoding).splitlines() if ln]
    assert all(ln.startswith("matrix_rows: ") for ln in lines)
    # TXT lines are "<product>: <record dict repr>" — eval the dict repr
    records = [eval(ln.split(": ", 1)[1]) for ln in lines]
    return [(int(r["id"]), r["name"]) for r in records]


def _read_dbunit(encoding: str) -> list[Row]:
    # Parsed from bytes on purpose: the XML declaration, not the test, must supply the encoding.
    root = ET.fromstring(_file("dbunit.xml").read_bytes())
    return [(int(row.get("id")), row.get("name")) for row in root if row.tag == "matrix_rows"]


def _read_fixed_width(encoding: str) -> list[Row]:
    # FCW sources are UTF-8 by contract; re-encode a copy so FileUtil can parse the written file.
    source = _file("fcw")
    utf8_copy = source.with_suffix(".utf8.fcw")
    utf8_copy.write_text(source.read_text(encoding=encoding), encoding="utf-8")
    return [(int(row["id"]), row["name"]) for row in FileUtil.read_fixed_width_to_dict_list(utf8_copy)]


_READERS = {
    EXPORTER_CSV: _read_csv,
    EXPORTER_JSON: _read_json,
    EXPORTER_XML: _read_xml,
    EXPORTER_XLSX: _read_xlsx,
    EXPORTER_TXT: _read_txt,
    EXPORTER_DBUNIT: _read_dbunit,
    EXPORTER_FIXED_WIDTH: _read_fixed_width,
}


@pytest.mark.parametrize(
    ("descriptor", "encoding"), [("matrix.xml", "utf-8"), ("matrix_cp1252.xml", "cp1252")], ids=["utf-8", "cp1252"]
)
def test_every_registered_buffered_exporter_round_trips(descriptor: str, encoding: str):
    # Self-extending: a new registry entry must add a matrix reader
    assert set(_READERS) == set(_BUFFERED_EXPORTERS), "matrix reader missing for a registered exporter"

    shutil.rmtree(_DIR / "output", ignore_errors=True)
    try:
        DataMimicTest(test_dir=_DIR, filename=descriptor).test_with_timer()
        for name, reader in _READERS.items():
            assert sorted(reader(encoding)) == _EXPECTED, f"{name} ({encoding})"
    finally:
        shutil.rmtree(_DIR / "output", ignore_errors=True)


def test_write_read_roundtrip_through_every_pagination_mode():
    """Write a JSON file, read it back as a source with cyclic/pageSize/distribution/unique.
    All asserts are seed-independent (wraparound order, permutation blocks, distinctness)."""
    shutil.rmtree(_DIR / "output", ignore_errors=True)
    try:
        engine = DataMimicTest(test_dir=_DIR, filename="roundtrip.xml", capture_test_result=True)
        engine.test_with_timer()
        result = engine.capture_result()

        src = list(range(1, 7))  # 6 source rows

        # ordered + cyclic: exact wraparound sequence across the 4 pages
        assert [r["id"] for r in result["ordered_cyclic"]] == src * 2 + [1, 2, 3]

        # random + cyclic: each source-length block of the global sequence is a permutation
        rand = [r["id"] for r in result["random_cyclic"]]
        assert len(rand) == 15
        assert sorted(rand[0:6]) == src and sorted(rand[6:12]) == src
        assert len(set(rand[12:15])) == 3

        # unique across pages: every source row exactly once
        assert sorted(r["id"] for r in result["unique_pages"]) == src

        # not cyclic: more requested than available -> stops at the 6 source rows, in order
        assert [r["id"] for r in result["exhausted"]] == src
    finally:
        shutil.rmtree(_DIR / "output", ignore_errors=True)


_SOURCE_CASES = {
    SourceFileFormat.CSV: "from_csv",
    SourceFileFormat.JSON: "from_json",
    SourceFileFormat.XML: "from_xml",
    SourceFileFormat.XLSX: "from_xlsx",
    SourceFileFormat.DBUNIT_XML: "from_dbunit",
    SourceFileFormat.FIXED_WIDTH: "from_fcw",
    SourceFileFormat.WEIGHTED_CSV: "from_wgt_csv",
    SourceFileFormat.WEIGHTED_ENTITY_CSV: "from_wgt_ent_csv",
}
_SOURCE_EXPECTED = {
    "from_wgt_csv": {"Müller-€", "Zoë"},
    "from_wgt_ent_csv": {"Göteborg", "São Paulo"},
}


@pytest.fixture(scope="module")
def source_matrix_result() -> dict[str, list[dict]]:
    shutil.rmtree(_DIR / "output", ignore_errors=True)
    try:
        engine = DataMimicTest(test_dir=_DIR, filename="source_matrix.xml", capture_test_result=True)
        engine.test_with_timer()
        return engine.capture_result()
    finally:
        shutil.rmtree(_DIR / "output", ignore_errors=True)


def test_source_matrix_covers_every_source_format():
    assert set(_SOURCE_CASES) == set(SourceFileFormat), "source matrix case missing for a SourceFileFormat"


@pytest.mark.parametrize("source_format", list(_SOURCE_CASES), ids=lambda f: f.name)
def test_every_source_format_reads_non_ascii(source_format: SourceFileFormat, source_matrix_result):
    product = _SOURCE_CASES[source_format]
    names = {row["name"] for row in source_matrix_result[product]}
    assert names == _SOURCE_EXPECTED.get(product, {f"Müller-€-{i}" for i in range(1, 4)}), product
