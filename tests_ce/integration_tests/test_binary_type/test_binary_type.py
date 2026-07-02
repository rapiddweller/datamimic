"""DSL-driven coverage for <key type="binary">: happy paths, determinism, casting,
error handling, and the export boundary (base64 in CSV/XLSX/JSON, raw bytes in a DB BLOB).
"""

import base64
import csv
import json
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from datamimic_ce.data_mimic_test import DataMimicTest

_DIR = Path(__file__).resolve().parent


def _run(filename: str) -> dict:
    engine = DataMimicTest(test_dir=_DIR, filename=filename, capture_test_result=True)
    engine.test_with_timer()
    return engine.capture_result()


def _clean():
    shutil.rmtree(_DIR / "output", ignore_errors=True)
    for f in _DIR.glob("db/*.sqlite"):
        f.unlink(missing_ok=True)


# ---------- happy paths ----------


def test_binary_lengths_and_types():
    rows = _run("binary_happy.xml")["g"]
    assert len(rows) == 8
    for r in rows:
        assert isinstance(r["default_bin"], bytes) and 1 <= len(r["default_bin"]) <= 16
        assert isinstance(r["exact8"], bytes) and len(r["exact8"]) == 8
        assert isinstance(r["ranged"], bytes) and 4 <= len(r["ranged"]) <= 10
        assert r["hexlen"] == 8  # sibling script can len() the bytes


def test_binary_seeded_reproducible_unseeded_not():
    first = [r["exact8"] for r in _run("binary_happy.xml")["g"]]
    second = [r["exact8"] for r in _run("binary_happy.xml")["g"]]
    assert first == second  # rngSeed -> byte-identical

    a = [r["b"] for r in _run("binary_unseeded.xml")["g"]]
    b = [r["b"] for r in _run("binary_unseeded.xml")["g"]]
    assert a != b  # unseeded -> different runs differ


def test_binary_casts_str_script_to_utf8_bytes():
    rows = _run("binary_cast.xml")["g"]
    assert rows[0]["from_str"] == b"hello"


# ---------- error handling ----------


def test_binary_min_greater_than_max_raises():
    with pytest.raises(Exception, match=r"min length 9 exceeds max length 4"):
        _run("binary_min_gt_max.xml")


def test_binary_negative_length_raises():
    with pytest.raises(Exception, match=r"lengths must be >= 0"):
        _run("binary_negative.xml")


def test_binary_cannot_cast_int():
    with pytest.raises(Exception, match=r"type='binary' cannot convert value of type 'int'"):
        _run("binary_cast_invalid.xml")


# ---------- export boundary ----------


def test_binary_exports_base64_to_files_and_raw_to_db():
    _clean()
    try:
        result = _run("binary_export.xml")
        payloads = [r["payload"] for r in result["blobs"]]
        assert all(isinstance(p, bytes) and len(p) == 6 for p in payloads)

        # DB round-trip: raw bytes survive the sqlite BLOB write and read back with the same length
        assert [r["len_back"] for r in result["back"]] == [6, 6, 6, 6, 6]
        assert sorted(r["payload"] for r in result["back"]) == sorted(payloads)

        out = _DIR / "output"
        csv_file = next(out.rglob("*.csv"))
        with csv_file.open() as f:
            csv_rows = list(csv.DictReader(f, delimiter="|"))
        assert [base64.b64decode(r["payload"]) for r in csv_rows] == payloads  # base64 text in csv

        xlsx_rows = list(load_workbook(next(out.rglob("*.xlsx"))).active.iter_rows(values_only=True))
        col = xlsx_rows[0].index("payload")
        assert [base64.b64decode(r[col]) for r in xlsx_rows[1:]] == payloads  # base64 text in xlsx

        json_rows = json.loads(next(out.rglob("*.json")).read_text())
        decoded = [base64.b64decode(r["payload"]) for r in json_rows]
        assert decoded == payloads  # serializer base64-encodes non-utf8 bytes
    finally:
        _clean()
