"""DSL-driven edge-case coverage for XLSX read and write:
empty / header-only / ragged / blank-header / invalid files, mixed types,
cyclic reads, nested + None + list write values, and multi-chunk volume.
"""

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from datamimic_ce.data_mimic_test import DataMimicTest


def _write_fixture(test_dir: Path, rows: list[list]):
    """Write rows (row 0 = header) to fixture.xlsx that the read descriptors point at."""
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    wb.save(test_dir / "fixture.xlsx")


def _write_raw_fixture(test_dir: Path, content: bytes):
    (test_dir / "fixture.xlsx").write_bytes(content)


def _run(filename: str, test_dir: Path, product: str = "rows") -> list[dict]:
    engine = DataMimicTest(test_dir=test_dir, filename=filename, capture_test_result=True)
    engine.test_with_timer()
    return engine.capture_result().get(product, [])  # 0-record product is absent from the result


def _find_xlsx(test_dir: Path) -> Path:
    files = [f for f in test_dir.rglob("*.xlsx") if f.name != "fixture.xlsx"]
    assert len(files) == 1, f"expected one output xlsx, got {files}"
    return files[0]


# ---------- READ edge cases ----------


def test_read_empty_file_yields_zero_rows(xlsx_test_dir: Path):
    _write_fixture(xlsx_test_dir, [])  # no rows at all
    assert _run("read_count_only.xml", xlsx_test_dir) == []


def test_read_header_only_yields_zero_rows(xlsx_test_dir: Path):
    _write_fixture(xlsx_test_dir, [["x", "y", "z"]])  # header, no data
    assert _run("read_count_only.xml", xlsx_test_dir) == []


def test_read_ragged_rows_pad_and_truncate(xlsx_test_dir: Path):
    # header x,y,z; one short row (2 cells), one long row (4 cells)
    _write_fixture(xlsx_test_dir, [["x", "y", "z"], [1, 2], [4, 5, 6, 7]])
    rows = _run("read_xyz.xml", xlsx_test_dir)
    assert len(rows) == 2
    # short row: missing 3rd cell -> None; long row: 4th cell ignored, no None-keyed column.
    assert (rows[0]["rx"], rows[0]["ry"], rows[0]["rz"]) == (1, 2, None)
    assert (rows[1]["rx"], rows[1]["ry"], rows[1]["rz"]) == (4, 5, 6)
    assert None not in rows[1]  # extra 4th cell did not create a None key


def test_read_blank_header_columns_are_dropped(xlsx_test_dir: Path):
    _write_fixture(xlsx_test_dir, [[None, "y", None], [1, 2, 3]])  # only 'y' is a real column
    rows = _run("read_count_only.xml", xlsx_test_dir)
    assert rows == [{"y": 2}]


def test_read_invalid_file_raises_clear_error(xlsx_test_dir: Path):
    _write_raw_fixture(xlsx_test_dir, b"this is plainly not a zip-based xlsx")
    with pytest.raises(Exception, match=r"Invalid XLSX file"):
        _run("read_count_only.xml", xlsx_test_dir)


def test_read_mixed_types_and_cyclic_transform(xlsx_test_dir: Path):
    # columns i(int) f(float) s(str) n(None); 3 source rows, cyclic to count=10
    _write_fixture(
        xlsx_test_dir,
        [
            ["i", "f", "s", "n"],
            [1, 1.5, "aa", None],
            [2, 2.5, "bb", "x"],
            [3, 3.5, "cc", None],
        ],
    )
    rows = _run("read_complex_cyclic.xml", xlsx_test_dir)
    assert len(rows) == 10  # cyclic wrap over 3 source rows
    # every emitted row is a correct transform of one of the 3 source rows
    expected = {1001: ("AA", 3.0, False), 1002: ("BB", 5.0, True), 1003: ("CC", 7.0, False)}
    for r in rows:
        assert (r["upper"], r["price_2"], r["has_note"]) == expected[r["id_plus"]]
    assert {r["id_plus"] for r in rows} == {1001, 1002, 1003}  # all sources seen
    assert len(rows) > len(expected)  # cyclic: rows repeat to reach count=10


# ---------- WRITE edge cases ----------


def test_write_none_nested_and_list_cells(xlsx_test_dir: Path):
    _run("write_edge.xml", xlsx_test_dir, product="edge")
    rows = list(load_workbook(_find_xlsx(xlsx_test_dir)).active.iter_rows(values_only=True))
    header = rows[0]
    assert "id" in header and "maybe_null" in header and "nested" in header
    body = rows[1:]
    assert len(body) == 3
    col = {name: i for i, name in enumerate(header)}
    assert body[0][col["maybe_null"]] is None  # None stays an empty cell
    assert body[0][col["nested"]] == '{"inner": "7"}'  # nested dict stringified as JSON


def test_write_multi_chunk_keeps_all_rows(xlsx_test_dir: Path):
    _run("write_chunked.xml", xlsx_test_dir, product="big")
    total = 0
    ids = set()
    for f in xlsx_test_dir.rglob("*.xlsx"):
        if f.name == "fixture.xlsx":
            continue
        body = list(load_workbook(f).active.iter_rows(values_only=True))[1:]
        total += len(body)
        ids.update(r[0] for r in body)
    assert total == 250  # all rows across chunk files
    assert ids == set(range(1, 251))  # every id present exactly once
